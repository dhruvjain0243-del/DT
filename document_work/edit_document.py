from docx import Document
from copy import deepcopy
src=r'C:\Users\dj845\Downloads\COA_PBL_II_Revised_Black_White.docx'
out=r'C:\Users\dj845\Desktop\DT\COA_PBL_II_Clean_Revised.docx'
d=Document(src)
changes={
9:'I would load the two operands from their fixed memory addresses into registers, then multiply them using register addressing. For example: LOAD R1, [1000]; LOAD R2, [1001]; MUL R1, R2. The multiplication can then use the register values without reading the operands from memory again.',
11:'I would choose RISC for this calculator. Simple instruction decoding can reduce the control hardware needed, which suits a low-cost design. The load-store approach also separates memory access from arithmetic and makes pipelining easier to organize.',
15:'The controller should handle the highest-priority pending interrupt first. With nested interrupts enabled, a more urgent request can interrupt a lower-priority service routine. The order shown above is a proposed order; the final priorities depend on the factory’s safety requirements.',
17:'Programmed I/O makes the CPU check the device and copy the data itself. For continuous vibration logs, I would use DMA. The CPU sets up the transfer, and the DMA controller copies the data to memory and signals completion. This gives the CPU more time to handle urgent sensor interrupts.',
20:'I would use memory-mapped I/O for this controller. The CPU can access the sensor registers with the same load and store instructions it uses for memory.',
24:'For example, MUL R1, R2, R3 followed by ADD R4, R1, R5 has a read-after-write (RAW) dependency. The ADD needs the value of R1 produced by the MUL, so it must wait if that result is not yet available.',
26:'Separate memory paths or additional hardware units can reduce structural hazards. Forwarding can pass a result directly to a dependent instruction, though some data hazards still require a stall. Branch prediction reduces the delay from control hazards when the prediction is correct. If it is wrong, the processor discards the instructions fetched along the wrong path.',
28:'Graphics processing applies many of the same operations to different vertices and pixels. A vector instruction such as VADD C, A, B adds corresponding elements of two vectors. It replaces several scalar additions with one instruction, reducing instruction overhead and improving throughput for suitable rendering operations.',
31:'Consider two caches, each with 16 lines. In a direct-mapped cache, line = block mod 16. For the sequence 0, 16, 32, 48, 0, 16, 32, 48, every block maps to line 0. Each access replaces the block already there, so all eight accesses miss. A four-way cache of the same capacity has four sets, with set = block mod 4. These blocks still map to set 0, but its four ways can hold all four blocks. The first four accesses miss; the next four hit. LRU selects a block to replace only when a new block needs space in a full set. This example shows how associativity reduces conflict misses.',
37:'I would use write-back for the server to reduce main-memory traffic. A write updates the cached line and marks it dirty. The changed data is written to lower-level memory when the line is evicted. Repeated writes to the same cached line therefore avoid repeated memory writes. Cache coherence and data protection still need attention because main memory may hold an older copy.',
40:'In a Cortex-M processor, arithmetic mainly uses registers, while load and store instructions access memory. Simple decoding and compact Thumb encodings help limit hardware and program size. These features suit a fitness tracker, where cost and battery life matter. Thumb instructions may be 16 or 32 bits, depending on the instruction and the core.',
42:'A three-stage pipeline can overlap instruction fetch (IF), decode (ID), and execution (EX). The example below shows three instructions moving through these stages, assuming no stalls:',
44:'A load-use hazard can occur with LDR R1, [R0] followed by ADD R2, R1, R3. The ADD depends on the value being loaded into R1 and may reach execution before that value is ready. Forwarding, instruction scheduling, or a stall can handle this, depending on the implementation. The table is a conceptual example; actual pipeline details vary between Cortex-M models.',
46:'For the tracker, I would use registers for current operands, on-chip SRAM for sensor buffers and working data, and Flash for firmware. A small cache may help if the selected Cortex-M core supports it and the workload benefits. A large cache could add area and power consumption with little benefit for a simple wearable.',
48:'These examples show why hardware choices depend on the workload. A calculator benefits from simple arithmetic hardware, while a factory controller needs timely interrupt handling. Graphics processing benefits from vector operations, and a busy server benefits from fewer memory accesses. For a fitness tracker, the design must also keep power use low.'
}
for i,text in changes.items():
 p=d.paragraphs[i]
 fmt=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear()
 r=p.add_run(text)
 if fmt is not None:r._r.insert(0,fmt)
for s in d.sections:
 for footer in (s.footer,s.first_page_footer,s.even_page_footer):
  for p in footer.paragraphs:p.clear()
d.save(out)
print(out)
print('Revised paragraphs:',len(changes),'Tables:',len(d.tables))
